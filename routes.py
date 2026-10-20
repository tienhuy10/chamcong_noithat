from flask import render_template, redirect, url_for, request, flash, jsonify, get_flashed_messages, session, send_file
from flask_login import login_user, logout_user, login_required, current_user
from models import db, User, Attendance, LeaveRequest, ShiftConfig, SupplementaryAttendance, ShareToken
from app import app, login_manager
from datetime import datetime, timedelta, time, date
import uuid
import os
import io
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
import pytz

vn_tz = pytz.timezone('Asia/Ho_Chi_Minh')

# Dictionary tạm để lưu token chia sẻ
shared_links = {}
share_tokens = {}

def calculate_work(attendance):
    morning_hours = 0.0
    afternoon_hours = 0.0

    # Chuẩn hóa về timezone 'Asia/Ho_Chi_Minh'
    def to_vn_time(dt):
        if dt is None:
            return None
        # Nếu datetime chưa aware, gán timezone VN
        if dt.tzinfo is None:
            return vn_tz.localize(dt)
        # Nếu đã aware nhưng không phải VN, chuyển về VN
        return dt.astimezone(vn_tz)

    morning_check_in = to_vn_time(attendance.morning_check_in)
    morning_check_out = to_vn_time(attendance.morning_check_out)
    afternoon_check_in = to_vn_time(attendance.afternoon_check_in)
    afternoon_check_out = to_vn_time(attendance.afternoon_check_out)

    if morning_check_in and morning_check_out:
        morning_hours = (morning_check_out - morning_check_in).total_seconds() / 3600

    if afternoon_check_in and afternoon_check_out:
        afternoon_hours = (afternoon_check_out - afternoon_check_in).total_seconds() / 3600

    attendance.morning_work_hours = morning_hours
    attendance.afternoon_work_hours = afternoon_hours
    attendance.work_hours = morning_hours + afternoon_hours


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password) and user.is_active:
            login_user(user)
            return redirect(url_for('employee_dashboard' if user.role == 'employee' else 'admin_dashboard'))
        flash('Invalid credentials or inactive account')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/employee/dashboard')
@login_required
def employee_dashboard():
    if current_user.role != 'employee':
        return redirect(url_for('admin_dashboard'))

    # Lấy ngày hiện tại và tuần hiện tại
    now = date.today()
    week_start = now - timedelta(days=now.weekday())  # Thứ 2
    week_end = week_start + timedelta(days=6)  # Chủ Nhật

    # Lịch sử chấm công (chỉ trong tuần)
    attendances = Attendance.query.filter_by(user_id=current_user.id).filter(
        Attendance.date >= week_start,
        Attendance.date <= week_end
    ).order_by(Attendance.date.desc()).all()

    # Tổng ngày làm
    total_work_days = sum(1 for att in attendances if att.work_hours > 0)

    # Chấm công bổ sung (chỉ trong tuần)
    supp_atts = SupplementaryAttendance.query.filter_by(user_id=current_user.id).filter(
        SupplementaryAttendance.date >= week_start,
        SupplementaryAttendance.date <= week_end
    ).order_by(SupplementaryAttendance.date.desc()).all()

    # Lịch nghỉ (chỉ pending và approved)
    leaves = LeaveRequest.query.filter_by(user_id=current_user.id).filter(
        LeaveRequest.date >= week_start,
        LeaveRequest.date <= week_end,
        LeaveRequest.status.in_(['pending', 'approved'])
    ).join(User).with_entities(
        LeaveRequest.date, LeaveRequest.shift, LeaveRequest.status, User.full_name.label('name')
    ).all()

    leave_dict = {}
    for leave in leaves:
        if leave.shift == 'full':
            for shift in ['morning', 'afternoon']:
                key = (leave.date, shift)
                if key not in leave_dict:
                    leave_dict[key] = []
                leave_dict[key].append({'name': leave.name, 'status': leave.status})
        else:
            key = (leave.date, leave.shift)
            if key not in leave_dict:
                leave_dict[key] = []
            leave_dict[key].append({'name': leave.name, 'status': leave.status})

    days = [week_start + timedelta(days=i) for i in range(7)]  # Thứ 2 đến Chủ Nhật

    return render_template(
        'employee_dashboard.html',
        attendances=attendances,
        total_work_days=total_work_days,
        supp_atts=supp_atts,
        leaves=leave_dict,
        days=days,
        now=now,
        week_start=week_start,
        week_end=week_end
    )

@app.route('/employee/attendance', methods=['GET', 'POST'])
@login_required
def attendance_form():
    today = date.today()
    current_time = datetime.now(vn_tz).time()
    attendance = Attendance.query.filter_by(user_id=current_user.id, date=today).first()

    # Kiểm tra trạng thái chấm công
    already_checked = {
        'morning': {'in': False, 'out': False},
        'afternoon': {'in': False, 'out': False}
    }
    if attendance:
        already_checked['morning']['in'] = bool(attendance.morning_check_in)
        already_checked['morning']['out'] = bool(attendance.morning_check_out)
        already_checked['afternoon']['in'] = bool(attendance.afternoon_check_in)
        already_checked['afternoon']['out'] = bool(attendance.afternoon_check_out)

    if request.method == 'POST':
        if not attendance:
            attendance = Attendance(user_id=current_user.id, date=today)
            db.session.add(attendance)

        is_morning = current_time < time(12, 0)
        is_afternoon = current_time >= time(13, 0)
        check_type = request.form.get('type')
        late_reason = request.form.get('late_reason')

        if is_morning:
            if check_type == 'in':
                if attendance and attendance.morning_check_in:
                    flash('Bạn đã chấm công ca sáng hôm nay!', 'error')
                    return redirect(url_for('employee_dashboard'))
                if current_time > time(7, 30) and not late_reason:
                    flash('Bạn đi muộn, vui lòng nhập lý do!', 'error')
                    return render_template('attendance_form.html', is_morning=is_morning, is_afternoon=is_afternoon,
                                         show_late_reason=True, already_checked=already_checked)
                attendance.morning_check_in = datetime.now(vn_tz)
                attendance.morning_shift = True
                if current_time > time(7, 30):
                    attendance.late = True
                    attendance.morning_late_minutes = int((datetime.combine(today, current_time) - datetime.combine(today, time(7, 30))).total_seconds() / 60)
                    attendance.morning_late_reason = late_reason
                flash('Check-in ca sáng thành công', 'success')
            elif check_type == 'out':
                if not attendance or not attendance.morning_check_in:
                    flash('Chưa check-in ca sáng, không thể check-out!', 'error')
                    return redirect(url_for('employee_dashboard'))
                if attendance.morning_check_out:
                    flash('Bạn đã chấm công ca sáng hôm nay!', 'error')
                    return redirect(url_for('employee_dashboard'))
                attendance.morning_check_out = datetime.now(vn_tz)
                flash('Check-out ca sáng thành công', 'success')
        elif is_afternoon:
            if check_type == 'in':
                if attendance and attendance.afternoon_check_in:
                    flash('Bạn đã chấm công ca chiều hôm nay!', 'error')
                    return redirect(url_for('employee_dashboard'))
                if current_time > time(13, 0) and not late_reason:
                    flash('Bạn đi muộn, vui lòng nhập lý do!', 'error')
                    return render_template('attendance_form.html', is_morning=is_morning, is_afternoon=is_afternoon,
                                         show_late_reason=True, already_checked=already_checked)
                attendance.afternoon_check_in = datetime.now(vn_tz)
                attendance.afternoon_shift = True
                if current_time > time(13, 0):
                    attendance.late = True
                    attendance.afternoon_late_minutes = int((datetime.combine(today, current_time) - datetime.combine(today, time(13, 0))).total_seconds() / 60)
                    attendance.afternoon_late_reason = late_reason
                flash('Check-in ca chiều thành công', 'success')
            elif check_type == 'out':
                if not attendance or not attendance.afternoon_check_in:
                    flash('Chưa check-in ca chiều, không thể check-out!', 'error')
                    return redirect(url_for('employee_dashboard'))
                if attendance.afternoon_check_out:
                    flash('Bạn đã chấm công ca chiều hôm nay!', 'error')
                    return redirect(url_for('employee_dashboard'))
                attendance.afternoon_check_out = datetime.now(vn_tz)
                flash('Check-out ca chiều thành công', 'success')
        else:
            flash('Ngoài giờ ca làm, không thể chấm công!', 'error')
            return redirect(url_for('employee_dashboard'))

        if 'image' in request.files and request.files['image'].filename:
            image = request.files['image']
            filename = f"{current_user.id}_{datetime.now(vn_tz).timestamp()}.jpg"
            os.makedirs('static/images', exist_ok=True)
            image.save(os.path.join('static/images', filename))
            attendance.image_proof = filename
        elif 'location' in request.form and request.form['location']:
            attendance.location = request.form['location']
            print(f"Saved location: {attendance.location}")

        calculate_work(attendance)
        db.session.commit()
        return redirect(url_for('employee_dashboard'))

    is_morning = current_time < time(12, 0)
    is_afternoon = current_time >= time(13, 0)
    show_late_reason = (is_morning and current_time > time(7, 30) and not already_checked['morning']['in']) or \
                       (is_afternoon and current_time > time(13, 0) and not already_checked['afternoon']['in'])
    return render_template('attendance_form.html', is_morning=is_morning, is_afternoon=is_afternoon,
                         show_late_reason=show_late_reason, already_checked=already_checked)

@app.route('/employee/supplementary_attendance', methods=['GET', 'POST'])
@login_required
def supplementary_attendance():
    if request.method == 'POST':
        date = datetime.strptime(request.form['date'], '%Y-%m-%d').date()
        shift = request.form['shift']
        start_time_str = request.form['start_time']
        end_time_str = request.form['end_time']
        reason = request.form['reason']

        # Convert to time objects
        try:
            start_time = datetime.strptime(start_time_str, '%H:%M').time()
            end_time = datetime.strptime(end_time_str, '%H:%M').time()
        except ValueError:
            flash('Định dạng giờ không hợp lệ. Vui lòng dùng HH:MM.', 'error')
            return redirect(url_for('supplementary_attendance'))

        # Calculate hours
        start_dt = datetime.combine(date, start_time)
        end_dt = datetime.combine(date, end_time)
        hours = (end_dt - start_dt).total_seconds() / 3600

        if hours <= 0:
            flash('Giờ ra phải sau giờ vào!', 'error')
            return redirect(url_for('supplementary_attendance'))

        if 'image' in request.files and request.files['image'].filename:
            image = request.files['image']
            filename = f"{current_user.id}_{datetime.now(vn_tz).timestamp()}.jpg"
            os.makedirs('static/images', exist_ok=True)
            image.save(os.path.join('static/images', filename))
        else:
            flash('Vui lòng cung cấp ảnh chấm công!', 'error')
            return redirect(url_for('supplementary_attendance'))

        supp_att = SupplementaryAttendance(
            user_id=current_user.id,
            date=date,
            shift=shift,
            start_time=start_time,
            end_time=end_time,
            hours=hours,
            reason=reason,
            image_proof=filename,
            status='pending'
        )
        db.session.add(supp_att)
        db.session.commit()
        flash('Yêu cầu chấm công bổ sung đã được gửi, chờ duyệt')
        return redirect(url_for('employee_dashboard'))

    return render_template('supplementary_attendance.html')

@app.route('/leave/schedule')
@login_required
def leave_schedule():
    today = datetime.now(vn_tz).date()
    week_start = today - timedelta(days=today.weekday())
    days = [week_start + timedelta(days=i) for i in range(7)]
    leaves = LeaveRequest.query.filter(LeaveRequest.date >= week_start, LeaveRequest.date < week_start + timedelta(days=7)).all()
    leave_dict = {}
    for leave in leaves:
        if leave.shift == 'full':
            for shift in ['morning', 'afternoon']:
                key = (leave.date, shift)
                if key not in leave_dict:
                    leave_dict[key] = []
                leave_dict[key].append({'name': leave.user.full_name, 'status': leave.status})
        else:
            key = (leave.date, leave.shift)
            if key not in leave_dict:
                leave_dict[key] = []
            leave_dict[key].append({'name': leave.user.full_name, 'status': leave.status})
    return render_template('leave_schedule.html', leaves=leave_dict, days=days)

@app.route('/leave/request', methods=['POST'])
@login_required
def leave_request():
    date = datetime.strptime(request.form['date'], '%Y-%m-%d').date()
    shift = request.form['shift']
    reason = request.form['reason']
    leave = LeaveRequest(user_id=current_user.id, date=date, shift=shift, reason=reason)
    db.session.add(leave)
    db.session.commit()
    flash('Yêu cầu nghỉ đã được gửi')
    return redirect(url_for('leave_schedule'))

@app.route('/admin/dashboard')
@login_required
def admin_dashboard():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))

    today = datetime.now(vn_tz).date()
    month = today.month
    year = today.year

    if month == 12:
        next_month = datetime(year + 1, 1, 1)
    else:
        next_month = datetime(year, month + 1, 1)
    days_in_month = (next_month - timedelta(days=1)).day

    # Tổng quan chấm công hôm nay
    users = User.query.filter_by(role='employee').all()
    total_employees = len(users)
    today_attendances = Attendance.query.filter_by(date=today).all()
    today_leaves = LeaveRequest.query.filter_by(date=today, status='approved').all()
    present = sum(1 for att in today_attendances if att.morning_shift or att.afternoon_shift)
    late = sum(1 for att in today_attendances if att.late)
    absent = total_employees - present
    approved_leaves = len(today_leaves)

    # Dữ liệu biểu đồ tròn
    pie_data = {
        'labels': ['Present', 'Absent', 'Late', 'Approved Leaves'],
        'data': [present, absent, late, approved_leaves]
    }

    # Yêu cầu chờ duyệt
    pending_supp_atts = SupplementaryAttendance.query.filter_by(status='pending').all()
    pending_leaves = LeaveRequest.query.filter_by(status='pending').all()

    # Thống kê tháng
    month_attendances = Attendance.query.filter(
        Attendance.date >= datetime(year, month, 1).date(),
        Attendance.date <= datetime(year, month, days_in_month).date()
    ).all()
    total_work_days = sum(1 for att in month_attendances if att.work_hours > 0)
    overtime_hours = sum(max(0, att.work_hours - 8) for att in month_attendances)
    month_leaves = LeaveRequest.query.filter(
        LeaveRequest.date >= datetime(year, month, 1).date(),
        LeaveRequest.date <= datetime(year, month, days_in_month).date(),
        LeaveRequest.status == 'approved'
    ).all()
    total_leave_days = len(month_leaves)

    # Dữ liệu biểu đồ đường
    days = [datetime(year, month, d + 1).date() for d in range(days_in_month)]
    attendance_trend = [0] * days_in_month
    for att in month_attendances:
        day_index = att.date.day - 1
        attendance_trend[day_index] += 0.5 * (int(att.morning_shift) + int(att.afternoon_shift))

    line_data = {
        'labels': [day.strftime('%d/%m') for day in days],
        'data': attendance_trend
    }

    # Lịch sử gần đây
    recent_supp_atts = SupplementaryAttendance.query.order_by(
        SupplementaryAttendance.date.desc(),
        SupplementaryAttendance.start_time.desc()
    ).limit(5).all()
    recent_leaves = LeaveRequest.query.order_by(LeaveRequest.date.desc()).limit(5).all()
    recent_actions = sorted(
        [(s, 'Supplementary Attendance', s.status) for s in recent_supp_atts] +
        [(l, 'Leave Request', l.status) for l in recent_leaves],
        key=lambda x: x[0].date if hasattr(x[0], 'date') else x[0].date,
        reverse=True
    )[:5]

    # Chấm công trong 1 tuần
    start_of_week = today - timedelta(days=today.weekday())
    end_of_week = start_of_week + timedelta(days=6)
    weekly_attendances = Attendance.query.filter(
        Attendance.date >= start_of_week,
        Attendance.date <= end_of_week
    ).order_by(Attendance.date.desc(), Attendance.user_id).all()

    # Lấy cấu hình giờ làm từ ShiftConfig
    shift_config = ShiftConfig.query.first()
    if shift_config:
        morning_threshold = shift_config.morning_start
    else:
        morning_threshold = time(7, 30)  # Giá trị mặc định nếu không có ShiftConfig

    # Tính toán lại late_minutes bằng cách trừ thời gian check-in với thời gian vào làm
    for att in weekly_attendances:
        if att.morning_check_in:
            check_in_datetime = att.morning_check_in.replace(tzinfo=vn_tz)
            threshold_datetime = datetime.combine(att.date, morning_threshold).replace(tzinfo=vn_tz)
            if check_in_datetime > threshold_datetime:
                late_minutes = int((check_in_datetime - threshold_datetime).total_seconds() / 60)
                if late_minutes > 0:
                    att.late = True
                    att.late_minutes = late_minutes
                    db.session.commit()
                else:
                    att.late = False
                    att.late_minutes = 0
                    db.session.commit()

    return render_template('admin_dashboard.html',
                         total_employees=total_employees, present=present, absent=absent, late=late,
                         approved_leaves=approved_leaves, pie_data=pie_data,
                         pending_supp_atts=pending_supp_atts, pending_leaves=pending_leaves,
                         total_work_days=total_work_days, overtime_hours=overtime_hours,
                         total_leave_days=total_leave_days, line_data=line_data,
                         recent_actions=recent_actions, month=month, year=year,
                         weekly_attendances=weekly_attendances)

@app.route('/admin/users', methods=['GET', 'POST'])
@login_required
def admin_users():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))
    users = User.query.all()
    if request.method == 'POST':
        if 'username' in request.form:  # New user addition
            username = request.form['username']
            password = request.form['password']
            full_name = request.form['full_name']
            user = User(username=username, full_name=full_name)
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash('User added')
        return redirect(url_for('admin_users'))
    return render_template('admin_users.html', users=users)

@app.route('/admin/user/<int:user_id>/edit', methods=['GET', 'POST'])
@login_required
def admin_edit_user(user_id):
    user = User.query.get_or_404(user_id)
    if request.method == 'POST':
        if 'activate' in request.form:
            user.is_active = True
            db.session.commit()
            flash('User activated')
        elif 'deactivate' in request.form:
            user.is_active = False
            db.session.commit()
            flash('User deactivated')
        elif 'delete' in request.form:
            if user.role == 'admin':
                flash('Cannot delete admin account', 'error')
            else:
                db.session.delete(user)
                db.session.commit()
                flash('User deleted')
        elif 'edit' in request.form:  # Edit user details
            username = request.form['username']
            full_name = request.form['full_name']
            password = request.form['password']
            if username:
                user.username = username
            if full_name:
                user.full_name = full_name
            if password:
                user.set_password(password)
            db.session.commit()
            flash('User updated')
        return redirect(url_for('admin_users'))
    return render_template('edit_user.html', user=user)

@app.route('/admin/attendance', methods=['GET', 'POST'])
@login_required
def admin_attendance():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))
    now = datetime.now(vn_tz)
    month = int(request.form.get('month', now.month))
    year = int(request.form.get('year', now.year))
    period = request.form.get('period', 'month')  # 'week' or 'month'

    if month == 12:
        next_month_first = datetime(year + 1, 1, 1)
    else:
        next_month_first = datetime(year, month + 1, 1)
    days_in_month = (next_month_first - timedelta(days=1)).day
    
    if period == 'week':
        # Get the first day of the week (Monday) containing the current date
        start_of_week = now - timedelta(days=now.weekday())
        days = [start_of_week + timedelta(days=i) for i in range(7) if start_of_week.month == month and start_of_week.year == year]
    else:  # month
        days = [datetime(year, month, d + 1).date() for d in range(days_in_month)]

    users = User.query.filter_by(role='employee').all()
    attendances = Attendance.query.filter(Attendance.date >= datetime(year, month, 1).date(),
                                         Attendance.date <= datetime(year, month, days_in_month).date()).all()
    leaves = LeaveRequest.query.filter(LeaveRequest.date >= datetime(year, month, 1).date(),
                                      LeaveRequest.date <= datetime(year, month, days_in_month).date()).all()
    supp_atts = SupplementaryAttendance.query.filter_by(status='pending').all()

    leave_dict = {(l.user_id, l.date, l.shift): l for l in leaves if l.shift != 'full'}
    leave_dict.update({(l.user_id, l.date, shift): l for l in leaves if l.shift == 'full' for shift in ['morning', 'afternoon']})
    attendances_dict = {(att.user_id, att.date): att for att in attendances}
    table = {user.id: {day: {'morning': '', 'afternoon': '', 'work': 0} for day in days} for user in users}

    for att in attendances:
        key = (att.user_id, att.date)
        morning_status = 'X' if att.morning_shift else 'Vắng'
        afternoon_status = 'X' if att.afternoon_shift else 'Vắng'
        if (att.user_id, att.date, 'morning') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'morning')]
            if leave.status == 'approved':
                morning_status = 'Nghỉ'
            elif leave.status == 'pending':
                morning_status = 'Nghỉ (Chờ)' if not att.morning_shift else 'X (Xin Nghỉ Chờ)'
        if (att.user_id, att.date, 'afternoon') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'afternoon')]
            if leave.status == 'approved':
                afternoon_status = 'Nghỉ'
            elif leave.status == 'pending':
                afternoon_status = 'Nghỉ (Chờ)' if not att.afternoon_shift else 'X (Xin Nghỉ Chờ)'
        if att.date in days:
            table[att.user_id][att.date] = {'morning': morning_status, 'afternoon': afternoon_status, 'work': 0}

    # Update table with approved supplementary attendance
    approved_supp_atts = SupplementaryAttendance.query.filter_by(status='approved').all()
    for supp_att in approved_supp_atts:
        att = Attendance.query.filter_by(user_id=supp_att.user_id, date=supp_att.date).first()
        if not att:
            att = Attendance(user_id=supp_att.user_id, date=supp_att.date)
            db.session.add(att)
        if supp_att.shift == 'morning':
            att.morning_shift = True
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['morning'] = 'X'
        elif supp_att.shift == 'afternoon':
            att.afternoon_shift = True
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['afternoon'] = 'X'
        db.session.commit()

    for leave in leaves:
        if leave.date in days and leave.date in table[leave.user_id]:
            if leave.status == 'approved':
                if leave.shift == 'full':
                    for shift in ['morning', 'afternoon']:
                        table[leave.user_id][leave.date][shift] = 'Nghỉ'
                elif leave.shift in ['morning', 'afternoon']:
                    table[leave.user_id][leave.date][leave.shift] = 'Nghỉ'
            elif leave.status == 'pending':
                if leave.shift == 'full':
                    for shift in ['morning', 'afternoon']:
                        table[leave.user_id][leave.date][shift] = 'Nghỉ (Chờ)'
                elif leave.shift in ['morning', 'afternoon']:
                    table[leave.user_id][leave.date][leave.shift] = 'Nghỉ (Chờ)'

    # Calculate work totals based on shifts
    work_totals = {}
    for user in users:
        total_work = 0
        for day in days:
            data = table[user.id][day]
            if data['morning'] == 'X' and data['afternoon'] == 'X':
                total_work += 1.0  # Full day (2 shifts)
            elif data['morning'] == 'X' or data['afternoon'] == 'X':
                total_work += 0.5  # Half day (1 shift)
        work_totals[user.id] = total_work

    total_work_days = sum(work_totals.values())
    leave_count = len([l for l in leaves if l.status == 'approved'])

    return render_template('admin_attendance.html', table=table, users=users, days=days, attendances=attendances_dict, work_totals=work_totals, leave_dict=leave_dict, supp_atts=supp_atts, total_work_days=total_work_days, leave_count=leave_count, month=month, year=year, period=period, datetime=datetime)

@app.route('/admin/edit_attendance/<int:att_id>', methods=['GET', 'POST'])
@login_required
def admin_edit_attendance(att_id):
    att = Attendance.query.get_or_404(att_id)
    if request.method == 'POST':
        att.morning_check_in = datetime.strptime(request.form['morning_check_in'], '%Y-%m-%dT%H:%M') if request.form['morning_check_in'] else None
        att.morning_check_out = datetime.strptime(request.form['morning_check_out'], '%Y-%m-%dT%H:%M') if request.form['morning_check_out'] else None
        att.afternoon_check_in = datetime.strptime(request.form['afternoon_check_in'], '%Y-%m-%dT%H:%M') if request.form['afternoon_check_in'] else None
        att.afternoon_check_out = datetime.strptime(request.form['afternoon_check_out'], '%Y-%m-%dT%H:%M') if request.form['afternoon_check_out'] else None
        att.late_reason = request.form.get('late_reason')
        att.morning_shift = bool(att.morning_check_in and att.morning_check_out)
        att.afternoon_shift = bool(att.afternoon_check_in and att.afternoon_check_out)
        calculate_work(att)
        db.session.commit()
        flash('Attendance updated')
        return redirect(url_for('attendance_detail', att_id=att_id))
    return render_template('admin_edit_attendance.html', att=att)

@app.route('/admin/attendance_detail/<int:att_id>', methods=['GET'])
@login_required
def attendance_detail(att_id):
    att = Attendance.query.get_or_404(att_id)
    leave = LeaveRequest.query.filter_by(user_id=att.user_id, date=att.date, status='approved').first()

    # Calculate total work time
    total_time = 0
    if att.morning_check_in and att.morning_check_out:
        total_time += (att.morning_check_out - att.morning_check_in).total_seconds() / 3600
    if att.afternoon_check_in and att.afternoon_check_out:
        total_time += (att.afternoon_check_out - att.afternoon_check_in).total_seconds() / 3600

    # Determine if late (example threshold: after 8:00 AM)
    is_late = False
    if att.morning_check_in and att.morning_check_in.time() > datetime.strptime('08:00', '%H:%M').time():
        is_late = True

    return render_template('attendance_detail.html', att=att, leave=leave, total_time=total_time, is_late=is_late)

# @app.route('/admin/share_attendance', methods=['POST'])
# @login_required
# def admin_share_attendance():
#     if current_user.role != 'admin':
#         return redirect(url_for('employee_dashboard'))
#     month = int(request.form['month'])
#     year = int(request.form['year'])
#     token = str(uuid.uuid4())
#     shared_links[token] = {'month': month, 'year': year}
#     share_url = url_for('shared_attendance', token=token, _external=True)
#     flash(f'Share link: {share_url}')
#     return redirect(url_for('admin_attendance'))

@app.route('/admin/share_attendance', methods=['POST'])
@login_required
def admin_share_attendance():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))

    month = int(request.form['month'])
    year = int(request.form['year'])
    period = request.form.get('period', 'month')

    share_token = str(uuid.uuid4())
    expiration_days = 10
    expiration_time = datetime.now(vn_tz) + timedelta(days=expiration_days)

    # Lưu vào database
    token_entry = ShareToken(
        token=share_token,
        month=month,
        year=year,
        period=period,
        expiration_time=expiration_time
    )
    db.session.add(token_entry)
    db.session.commit()

    share_url = url_for('shared_attendance', token=share_token, _external=True)
    flash(f'Link chia sẻ: {share_url} (hết hạn sau {expiration_days} ngày)')
    return redirect(url_for('admin_attendance'))



@app.route('/share/<token>')
def shared_attendance(token):
    share = ShareToken.query.filter_by(token=token).first()
    if not share or share.expiration_time < datetime.now():
        return "❌ Link đã hết hạn hoặc không hợp lệ.", 404

    month = share.month
    year = share.year
    period = share.period

    days_in_month = (datetime(year, month + 1, 1) - timedelta(days=1)).day
    if period == 'week':
        start_of_week = datetime.now(vn_tz) - timedelta(days=datetime.now(vn_tz).weekday())
        days = [start_of_week + timedelta(days=i) for i in range(7) if start_of_week.month == month and start_of_week.year == year]
    else:  # month
        days = [datetime(year, month, d + 1).date() for d in range(days_in_month)]

    # Chỉ lấy các nhân viên đang hoạt động (is_active=True)
    users = User.query.filter_by(role='employee', is_active=True).all()
    attendances = Attendance.query.filter(Attendance.date >= datetime(year, month, 1).date(),
                                         Attendance.date <= datetime(year, month, days_in_month).date()).all()
    leaves = LeaveRequest.query.filter(LeaveRequest.date >= datetime(year, month, 1).date(),
                                      LeaveRequest.date <= datetime(year, month, days_in_month).date()).all()

    leave_dict = {(l.user_id, l.date, l.shift): l for l in leaves if l.shift != 'full'}
    leave_dict.update({(l.user_id, l.date, shift): l for l in leaves if l.shift == 'full' for shift in ['morning', 'afternoon']})
    attendances_dict = {(att.user_id, att.date): att for att in attendances}
    table = {user.id: {day: {'morning': '', 'afternoon': '', 'work': 0} for day in days} for user in users}

    for att in attendances:
        # Chỉ xử lý bản ghi chấm công của người dùng đang hoạt động
        if att.user_id not in table:
            continue
        key = (att.user_id, att.date)
        morning_status = 'X' if att.morning_shift else 'Vắng'
        afternoon_status = 'X' if att.afternoon_shift else 'Vắng'
        if (att.user_id, att.date, 'morning') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'morning')]
            if leave.status == 'approved':
                morning_status = 'Nghỉ'
        if (att.user_id, att.date, 'afternoon') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'afternoon')]
            if leave.status == 'approved':
                afternoon_status = 'Nghỉ'
        if att.date in days:
            table[att.user_id][att.date] = {'morning': morning_status, 'afternoon': afternoon_status, 'work': 0}

    approved_supp_atts = SupplementaryAttendance.query.filter_by(status='approved').all()
    for supp_att in approved_supp_atts:
        # Chỉ xử lý chấm công bổ sung của người dùng đang hoạt động
        if supp_att.user_id not in table:
            continue
        att = Attendance.query.filter_by(user_id=supp_att.user_id, date=supp_att.date).first()
        if not att:
            att = Attendance(user_id=supp_att.user_id, date=supp_att.date)
            db.session.add(att)
        if supp_att.shift == 'morning':
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['morning'] = 'X'
        elif supp_att.shift == 'afternoon':
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['afternoon'] = 'X'
        db.session.commit()

    work_totals = {}
    for user in users:
        total_work = 0
        for day in days:
            data = table[user.id][day]
            if data['morning'] == 'X' and data['afternoon'] == 'X':
                total_work += 1.0
            elif data['morning'] == 'X' or data['afternoon'] == 'X':
                total_work += 0.5
        work_totals[user.id] = total_work

    if request.args.get('export') == 'excel':
        # Create a new workbook and select the active sheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Title row
        ws['A1'] = f"Bảng Chấm Công {'Tháng' if period == 'month' else 'Tuần'} {month:02d}/{year}"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(users) + 1)
        ws.row_dimensions[1].height = 20
        ws['A1'].font = Font(bold=True, size=14)

        # Empty row
        ws['A2'] = ""
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(users) + 1)

        # Header row
        ws['A3'] = "Ngày"
        for col, user in enumerate(users, start=2):
            ws.cell(row=3, column=col, value=user.full_name if user.full_name else f'User_{user.id}')
        ws.row_dimensions[3].font = Font(bold=True)

        # Data rows
        for row, day in enumerate(days, start=4):
            ws.cell(row=row, column=1, value=day.strftime('%d/%m/%Y'))
            for col, user in enumerate(users, start=2):
                data = table[user.id][day]
                if data['morning'] == 'X' and data['afternoon'] == 'X':
                    status = '✓'
                    fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                elif data['morning'] == 'X' and data['afternoon'] != 'X':
                    status = 'Nghỉ Chiều'
                    fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
                elif data['morning'] != 'X' and data['afternoon'] == 'X':
                    status = 'Nghỉ Sáng'
                    fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
                else:
                    status = 'Nghỉ'
                    fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                ws.cell(row=row, column=col, value=status).fill = fill

        # Total row
        ws.cell(row=row + 1, column=1, value="Tổng Ngày Công")
        ws.row_dimensions[row + 1].font = Font(bold=True)
        for col, user in enumerate(users, start=2):
            ws.cell(row=row + 1, column=col, value=f"{work_totals.get(user.id, 0):.1f}")
            ws.cell(row=row + 1, column=col).fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")

        # Adjust column widths
        max_col = len(users) + 1  # Total number of columns (Ngày + users)
        for col in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15

        # Save to BytesIO
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'attendance_{datetime.now(vn_tz).strftime("%Y-%m-%d_%H-%M")}.xlsx'
        return send_file(output, download_name=filename, as_attachment=True)

    return render_template('shared_attendance.html', table=table, users=users, days=days, work_totals=work_totals, share=share)


@app.route('/admin/stats')
@login_required
def admin_stats():
    now = datetime.now(vn_tz)
    month = now.month
    year = now.year
    if month == 12:
        next_month_first = datetime(year + 1, 1, 1)
    else:
        next_month_first = datetime(year, month + 1, 1)

    days_in_month = (next_month_first - timedelta(days=1)).day
    start = datetime(year, month, 1).date()
    end = datetime(year, month, days_in_month).date()

    periods = {
        '3_days': (now.date() - timedelta(days=2), now.date()),
        'week': (now.date() - timedelta(days=now.weekday()), now.date()),
        'month': (start, end)
    }
    stats = {}
    for period, (start, end) in periods.items():
        attendances = Attendance.query.filter(Attendance.date >= start, Attendance.date <= end).all()
        late_count = sum(1 for a in attendances if a.late)
        early_count = sum(1 for a in attendances if a.early_leave)
        missing_hours = sum(max(8 - a.work_hours, 0) for a in attendances)
        work_total = sum(0.5 * (int(a.morning_shift) + int(a.afternoon_shift)) for a in attendances)
        stats[period] = {'late': late_count, 'early': early_count, 'missing': missing_hours, 'work': work_total}

    leaves = LeaveRequest.query.filter(LeaveRequest.date >= start, LeaveRequest.date <= end).all()
    leave_reasons = [(l.user.full_name, l.date, l.shift, l.reason, l.status) for l in leaves]

    return render_template('admin_stats.html', stats=stats, leave_reasons=leave_reasons, now=now)

# @app.route('/admin/edit_attendance/<int:att_id>', methods=['GET', 'POST'])
# @login_required
# def admin_edit_attendance(att_id):
#     att = Attendance.query.get(att_id)
#     if request.method == 'POST':
#         att.morning_check_in = datetime.strptime(request.form['morning_check_in'], '%Y-%m-%d %H:%M') if request.form['morning_check_in'] else None
#         att.morning_check_out = datetime.strptime(request.form['morning_check_out'], '%Y-%m-%d %H:%M') if request.form['morning_check_out'] else None
#         att.afternoon_check_in = datetime.strptime(request.form['afternoon_check_in'], '%Y-%m-%d %H:%M') if request.form['afternoon_check_in'] else None
#         att.afternoon_check_out = datetime.strptime(request.form['afternoon_check_out'], '%Y-%m-%d %H:%M') if request.form['afternoon_check_out'] else None
#         att.late_reason = request.form.get('late_reason')
#         calculate_work(att)
#         db.session.commit()
#         flash('Attendance updated')
#     return render_template('admin_edit_attendance.html', att=att)

@app.route('/admin/supplementary_attendances', methods=['GET'])
@login_required
def admin_supplementary_attendances():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))
    now = datetime.now(vn_tz)
    month = int(request.form.get('month', now.month))
    year = int(request.form.get('year', now.year))
    period = request.form.get('period', 'month')  # 'week' or 'month'

    if month == 12:
        next_month_first = datetime(year + 1, 1, 1)
    else:
        next_month_first = datetime(year, month + 1, 1)
    days_in_month = (next_month_first - timedelta(days=1)).day
    if period == 'week':
        # Get the first day of the week (Monday) containing the current date
        start_of_week = now - timedelta(days=now.weekday())
        days = [start_of_week + timedelta(days=i) for i in range(7) if start_of_week.month == month and start_of_week.year == year]
    else:  # month
        days = [datetime(year, month, d + 1).date() for d in range(days_in_month)]

    users = User.query.filter_by(role='employee').all()
    attendances = Attendance.query.filter(Attendance.date >= datetime(year, month, 1).date(),
                                         Attendance.date <= datetime(year, month, days_in_month).date()).all()
    leaves = LeaveRequest.query.filter(LeaveRequest.date >= datetime(year, month, 1).date(),
                                      LeaveRequest.date <= datetime(year, month, days_in_month).date()).all()
    supp_atts = SupplementaryAttendance.query.filter_by(status='pending').all()

    leave_dict = {(l.user_id, l.date, l.shift): l for l in leaves if l.shift != 'full'}
    leave_dict.update({(l.user_id, l.date, shift): l for l in leaves if l.shift == 'full' for shift in ['morning', 'afternoon']})
    attendances_dict = {(att.user_id, att.date): att for att in attendances}
    table = {user.id: {day: {'morning': '', 'afternoon': '', 'work': 0} for day in days} for user in users}

    for att in attendances:
        key = (att.user_id, att.date)
        morning_status = 'X' if att.morning_shift else 'Vắng'
        afternoon_status = 'X' if att.afternoon_shift else 'Vắng'
        if (att.user_id, att.date, 'morning') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'morning')]
            if leave.status == 'approved':
                morning_status = 'Nghỉ'
            elif leave.status == 'pending':
                morning_status = 'Nghỉ (Chờ)' if not att.morning_shift else 'X (Xin Nghỉ Chờ)'
        if (att.user_id, att.date, 'afternoon') in leave_dict:
            leave = leave_dict[(att.user_id, att.date, 'afternoon')]
            if leave.status == 'approved':
                afternoon_status = 'Nghỉ'
            elif leave.status == 'pending':
                afternoon_status = 'Nghỉ (Chờ)' if not att.afternoon_shift else 'X (Xin Nghỉ Chờ)'
        if att.date in days:
            table[att.user_id][att.date] = {'morning': morning_status, 'afternoon': afternoon_status, 'work': 0}

    # Update table with approved supplementary attendance
    approved_supp_atts = SupplementaryAttendance.query.filter_by(status='approved').all()
    for supp_att in approved_supp_atts:
        att = Attendance.query.filter_by(user_id=supp_att.user_id, date=supp_att.date).first()
        if not att:
            att = Attendance(user_id=supp_att.user_id, date=supp_att.date)
            db.session.add(att)
        if supp_att.shift == 'morning':
            att.morning_shift = True
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['morning'] = 'X'
        elif supp_att.shift == 'afternoon':
            att.afternoon_shift = True
            if supp_att.date in days:
                table[supp_att.user_id][supp_att.date]['afternoon'] = 'X'
        db.session.commit()

    for leave in leaves:
        if leave.date in days and leave.date in table[leave.user_id]:
            if leave.status == 'approved':
                if leave.shift == 'full':
                    for shift in ['morning', 'afternoon']:
                        table[leave.user_id][leave.date][shift] = 'Nghỉ'
                elif leave.shift in ['morning', 'afternoon']:
                    table[leave.user_id][leave.date][leave.shift] = 'Nghỉ'
            elif leave.status == 'pending':
                if leave.shift == 'full':
                    for shift in ['morning', 'afternoon']:
                        table[leave.user_id][leave.date][shift] = 'Nghỉ (Chờ)'
                elif leave.shift in ['morning', 'afternoon']:
                    table[leave.user_id][leave.date][leave.shift] = 'Nghỉ (Chờ)'

    # Calculate work totals based on shifts
    work_totals = {}
    for user in users:
        total_work = 0
        for day in days:
            data = table[user.id][day]
            if data['morning'] == 'X' and data['afternoon'] == 'X':
                total_work += 1.0  # Full day (2 shifts)
            elif data['morning'] == 'X' or data['afternoon'] == 'X':
                total_work += 0.5  # Half day (1 shift)
        work_totals[user.id] = total_work

    total_work_days = sum(work_totals.values())
    leave_count = len([l for l in leaves if l.status == 'approved'])

    return render_template('admin_supplementary_attendances.html', table=table, users=users, days=days, attendances=attendances_dict, work_totals=work_totals, leave_dict=leave_dict, supp_atts=supp_atts, total_work_days=total_work_days, leave_count=leave_count, month=month, year=year, period=period, datetime=datetime)

@app.route('/admin/supplementary_attendance/<int:supp_id>/approve', methods=['POST'])
@login_required
def approve_supplementary_attendance(supp_id):
    supp_att = SupplementaryAttendance.query.get(supp_id)
    if supp_att.status != 'pending':
        flash('Yêu cầu đã được xử lý!')
        return redirect(url_for('admin_dashboard'))

    attendance = Attendance.query.filter_by(user_id=supp_att.user_id, date=supp_att.date).first()
    if not attendance:
        attendance = Attendance(user_id=supp_att.user_id, date=supp_att.date)
        db.session.add(attendance)

    if supp_att.shift == 'morning':
        attendance.morning_check_in = datetime.combine(supp_att.date, supp_att.start_time)
        attendance.morning_check_out = datetime.combine(supp_att.date, supp_att.end_time) if supp_att.end_time else None
        if supp_att.start_time > time(7, 30):
            attendance.late = True
            attendance.late_reason = supp_att.reason
    elif supp_att.shift == 'afternoon':
        attendance.afternoon_check_in = datetime.combine(supp_att.date, supp_att.start_time)
        attendance.afternoon_check_out = datetime.combine(supp_att.date, supp_att.end_time) if supp_att.end_time else None

    attendance.image_proof = supp_att.image_proof
    calculate_work(attendance)
    supp_att.status = 'approved'
    supp_att.admin_response = request.form.get('admin_response', '')
    db.session.commit()
    flash('Yêu cầu chấm công bổ sung đã được duyệt')
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/supplementary_attendance/<int:supp_id>/reject', methods=['POST'])
@login_required
def reject_supplementary_attendance(supp_id):
    supp_att = SupplementaryAttendance.query.get(supp_id)
    supp_att.status = 'rejected'
    supp_att.admin_response = request.form.get('admin_response', '')
    db.session.commit()
    flash('Yêu cầu chấm công bổ sung đã bị từ chối')
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/leaves')
@login_required
def admin_leaves():
    leaves = LeaveRequest.query.filter_by(status='pending').all()
    today = datetime.now(vn_tz).date()
    week_start = today - timedelta(days=today.weekday())
    days = [week_start + timedelta(days=i) for i in range(7)]
    all_leaves = LeaveRequest.query.filter(LeaveRequest.date >= week_start, LeaveRequest.date < week_start + timedelta(days=7)).all()
    leave_dict = {}
    for leave in all_leaves:
        if leave.shift == 'full':
            for shift in ['morning', 'afternoon']:
                key = (leave.date, shift)
                if key not in leave_dict:
                    leave_dict[key] = []
                leave_dict[key].append({'name': leave.user.full_name, 'status': leave.status})
        else:
            key = (leave.date, leave.shift)
            if key not in leave_dict:
                leave_dict[key] = []
            leave_dict[key].append({'name': leave.user.full_name, 'status': leave.status})
    return render_template('admin_leaves.html', leaves=leaves, leave_dict=leave_dict, days=days)

@app.route('/admin/leave/<int:leave_id>/approve', methods=['POST'])
@login_required
def approve_leave(leave_id):
    leave = LeaveRequest.query.get(leave_id)
    leave.status = 'approved'
    db.session.commit()
    flash('Yêu cầu nghỉ đã được duyệt')
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/leave/<int:leave_id>/reject', methods=['POST'])
@login_required
def reject_leave(leave_id):
    leave = LeaveRequest.query.get(leave_id)
    leave.status = 'rejected'
    db.session.commit()
    flash('Yêu cầu nghỉ đã bị từ chối')
    return redirect(url_for('admin_dashboard'))

# @app.route('/shared/attendance/<token>')
# def shared_attendance(token):
#     if token not in shared_links:
#         return 'Invalid link', 404
#     data = shared_links[token]
#     month = data['month']
#     year = data['year']
#     days_in_month = (datetime(year, month+1, 1) - timedelta(days=1)).day
#     days = [datetime(year, month, d+1).date() for d in range(days_in_month)]
#     users = User.query.filter_by(role='employee').all()
#     attendances = Attendance.query.filter(Attendance.date >= datetime(year, month, 1).date(),
#                                          Attendance.date <= datetime(year, month, days_in_month).date()).all()
#     table = {user.id: {day: {'morning': '', 'afternoon': ''} for day in days} for user in users}
#     for att in attendances:
#         table[att.user_id][att.date]['morning'] = 'X' if att.morning_shift else 'Vắng'
#         table[att.user_id][att.date]['afternoon'] = 'X' if att.afternoon_shift else 'Vắng'
#     return render_template('shared_attendance.html', table=table, users=users, days=days)

@app.route('/')
def index():
    if current_user.is_authenticated:
        if current_user.role == 'admin':
            return redirect(url_for('admin_dashboard'))
        return redirect(url_for('employee_dashboard'))
    return redirect(url_for('login'))

@app.route('/admin/shift_config', methods=['GET', 'POST'])
@login_required
def admin_shift_config():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))
    config = ShiftConfig.query.first()
    if not config:
        config = ShiftConfig()
        db.session.add(config)
        db.session.commit()

    if request.method == 'POST':
        config.morning_start = datetime.strptime(request.form['morning_start'], '%H:%M').time()
        config.morning_end = datetime.strptime(request.form['morning_end'], '%H:%M').time()
        config.afternoon_min = datetime.strptime(request.form['afternoon_min'], '%H:%M').time()
        config.afternoon_max = datetime.strptime(request.form['afternoon_max'], '%H:%M').time()
        config.afternoon_hours = int(request.form['afternoon_hours'])
        db.session.commit()
        flash('Thiết lập thời gian ca làm đã cập nhật')

    return render_template('admin_shift_config.html', config=config)

@app.route('/favicon.ico')
def favicon():
    return app.send_static_file('favicon.ico')

@app.route('/admin/data_management', methods=['GET', 'POST'])
@login_required
def admin_data_management():
    if current_user.role != 'admin':
        return redirect(url_for('employee_dashboard'))

    now = datetime.now(vn_tz)
    month = int(request.form.get('month', now.month))
    year = int(request.form.get('year', now.year))
    period = request.form.get('period', 'all')  # 'all', 'month', 'week', 'year'

    # Filter attendances based on period
    if period == 'month':
        start_date = datetime(year, month, 1).date()
        end_date = (datetime(year, month + 1, 1) - timedelta(days=1)).date()
        attendances = Attendance.query.filter(Attendance.date >= start_date, Attendance.date <= end_date).all()
    elif period == 'week':
        start_of_week = now - timedelta(days=now.weekday())
        end_of_week = start_of_week + timedelta(days=6)
        attendances = Attendance.query.filter(Attendance.date >= start_of_week.date(), Attendance.date <= end_of_week.date()).all()
    elif period == 'year':
        start_date = datetime(year, 1, 1).date()
        end_date = datetime(year, 12, 31).date()
        attendances = Attendance.query.filter(Attendance.date >= start_date, Attendance.date <= end_date).all()
    else:  # all
        attendances = Attendance.query.all()

    leave_requests = LeaveRequest.query.all()
    supp_atts = SupplementaryAttendance.query.all()  # Assuming late reasons are in SupplementaryAttendance

    if request.method == 'POST':
        if 'delete_attendance' in request.form:
            att_id = request.form['delete_attendance']
            att = Attendance.query.get(att_id)
            if att:
                db.session.delete(att)
                db.session.commit()
                flash('Chấm công đã được xóa')
        elif 'delete_leave' in request.form:
            leave_id = request.form['delete_leave']
            leave = LeaveRequest.query.get(leave_id)
            if leave:
                db.session.delete(leave)
                db.session.commit()
                flash('Đơn xin đã được xóa')
        elif 'delete_supp' in request.form:
            supp_id = request.form['delete_supp']
            supp = SupplementaryAttendance.query.get(supp_id)
            if supp:
                db.session.delete(supp)
                db.session.commit()
                flash('Lý do bổ sung đã được xóa')
        elif 'batch_delete_attendance' in request.form:
            if period == 'month':
                start_date = datetime(year, month, 1).date()
                end_date = (datetime(year, month + 1, 1) - timedelta(days=1)).date()
                db.session.query(Attendance).filter(Attendance.date >= start_date, Attendance.date <= end_date).delete()
            elif period == 'week':
                start_of_week = now - timedelta(days=now.weekday())
                end_of_week = start_of_week + timedelta(days=6)
                db.session.query(Attendance).filter(Attendance.date >= start_of_week.date(), Attendance.date <= end_of_week.date()).delete()
            elif period == 'year':
                start_date = datetime(year, 1, 1).date()
                end_date = datetime(year, 12, 31).date()
                db.session.query(Attendance).filter(Attendance.date >= start_date, Attendance.date <= end_date).delete()
            else:  # all
                db.session.query(Attendance).delete()
            db.session.commit()
            flash('Chấm công đã được xóa theo bộ lọc')
        elif 'batch_delete_leave' in request.form:
            db.session.query(LeaveRequest).delete()
            db.session.commit()
            flash('Tất cả đơn xin đã được xóa')
        elif 'batch_delete_supp' in request.form:
            db.session.query(SupplementaryAttendance).delete()
            db.session.commit()
            flash('Tất cả lý do bổ sung đã được xóa')
        return redirect(url_for('admin_data_management'))

    return render_template('admin_data_management.html', attendances=attendances, leave_requests=leave_requests, supp_atts=supp_atts, month=month, year=year, period=period)

